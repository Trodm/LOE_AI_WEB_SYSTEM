
from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path
from typing import Any, Dict, List, Optional

PDF_EXTENSIONS = {".pdf"}
WORD_EXTENSIONS = {".docx", ".doc"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
SUPPORTED_REPORT_EXTENSIONS = PDF_EXTENSIONS | WORD_EXTENSIONS | IMAGE_EXTENSIONS


def _norm_spaces(text: str) -> str:
    text = (text or "").replace("\x00", " ")
    text = re.sub(r"[\u00a0\u200b]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _read_pdf_text(path: Path) -> str:
    try:
        import pdfplumber
        parts = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                txt = page.extract_text() or ""
                if txt:
                    parts.append(txt)
        return "\n".join(parts)
    except Exception:
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            return ""


def _ocr_image(path: Path) -> str:
    try:
        from PIL import Image
        import pytesseract
        return pytesseract.image_to_string(Image.open(path)) or ""
    except Exception as exc:
        raise RuntimeError("OCR failed. Install Tesseract OCR locally or upload a searchable PDF/DOCX. Details: " + str(exc))


def _ocr_pdf(path: Path) -> str:
    try:
        import pdfplumber
        import pytesseract
        parts = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                image = page.to_image(resolution=200).original
                parts.append(pytesseract.image_to_string(image) or "")
        return "\n".join(parts)
    except Exception as exc:
        raise RuntimeError("Scanned PDF OCR failed. Install Tesseract OCR or upload a clearer/searchable PDF. Details: " + str(exc))


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("python-docx is required to read DOCX files.") from exc
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        if p.text:
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            vals = [c.text for c in row.cells if c.text]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def _read_doc(path: Path) -> str:
    try:
        import win32com.client  # type: ignore
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(path))
        text = doc.Content.Text
        doc.Close(False)
        word.Quit()
        return text or ""
    except Exception:
        antiword = shutil.which("antiword")
        if antiword:
            completed = subprocess.run([antiword, str(path)], capture_output=True, text=True, check=False)
            return completed.stdout or ""
        raise RuntimeError("Could not read .doc file. Save as .docx or run on Windows with Microsoft Word installed.")


def extract_ip_report_text(report_path: Path) -> Dict[str, Any]:
    report_path = Path(report_path)
    ext = report_path.suffix.lower()
    if ext not in SUPPORTED_REPORT_EXTENSIONS:
        raise ValueError("Unsupported IP report format. Upload PDF, scanned PDF, Word DOCX/DOC, or image files.")
    method = ""
    text = ""
    if ext == ".pdf":
        text = _read_pdf_text(report_path)
        method = "pdf-text"
        if len(_norm_spaces(text)) < 150:
            text = _ocr_pdf(report_path)
            method = "pdf-ocr"
    elif ext == ".docx":
        text = _read_docx(report_path); method = "docx-text"
    elif ext == ".doc":
        text = _read_doc(report_path); method = "doc-text"
    else:
        text = _ocr_image(report_path); method = "image-ocr"
    clean = _norm_spaces(text)
    if len(clean) < 50:
        raise RuntimeError("The system could not extract enough readable text from this IP report.")
    return {"text": text, "clean": clean, "method": method, "characters": len(clean)}


def _all_matches(patterns: List[str], clean: str) -> List[re.Match]:
    matches = []
    for p in patterns:
        matches.extend(re.finditer(p, clean, flags=re.I))
    return sorted(matches, key=lambda m: m.start())


def _extract_between_headings(clean: str, start_patterns: List[str], stop_patterns: List[str]) -> str:
    candidates = []
    for m in _all_matches(start_patterns, clean):
        start = m.start()
        end = len(clean)
        after = clean[start + len(m.group(0)):]
        for stop_pat in stop_patterns:
            sm = re.search(stop_pat, after, flags=re.I)
            if sm:
                end = min(end, start + len(m.group(0)) + sm.start())
        cand = _norm_spaces(clean[start:end])
        if "contents" in cand[:400].lower() or "......." in cand[:300]:
            score = len(cand) - 2000
        else:
            score = len(cand)
        if len(cand) > 100:
            candidates.append((score, start, cand))
    if not candidates:
        return ""
    # Prefer the longest true body section; table of contents entries are penalised.
    return sorted(candidates, key=lambda x: (x[0], x[1]), reverse=True)[0][2]


def extract_earnings_sections(clean: str) -> Dict[str, str]:
    pre = _extract_between_headings(
        clean,
        [
            r"\bPRE[-\s]?ACCIDENT\s+EMPLOYMENT\s*&\s*EARNINGS\s+SCENARIO\b",
            r"\bPRE[-\s]?ACCIDENT\s+EMPLOYMENT\s+SCENARIO\b",
            r"\bPRE[-\s]?MORBID\s+EARNINGS\b",
            r"\bPRE[-\s]?MORBID\s+POTENTIAL\s+FUNCTIONING\b",
            r"\bPRE[-\s]?ACCIDENT\s+SCENARIO\s+LOSS\s+OF\s+EMPLOYMENT\s+AND\s+EARNINGS\b",
        ],
        [
            r"\bPOST[-\s]?ACCIDENT\b", r"\bPOST[-\s]?MORBID\b", r"\bLOSS\s+OF\s+EARNINGS\b", r"\bSECTION\s+16\b", r"\b7\.2\b"
        ],
    )
    post = _extract_between_headings(
        clean,
        [
            r"\bPOST[-\s]?ACCIDENT\s+LOSS\s+OF\s+EMPLOYMENT\s+AND\s+EARNINGS\s+SCENARIO\b",
            r"\bPOST[-\s]?ACCIDENT\s+EMPLOYMENT\s+SCENARIO\b",
            r"\bPOST[-\s]?MORBID\s+PROBABLE\s+EARNING\s+PROJECTIONS\b",
            r"\bPOST[-\s]?MORBID\s+POTENTIAL\s+FUNCTIONING\b",
            r"\bLOSS\s+OF\s+EARNINGS\b",
        ],
        [r"\bRECOMMENDATIONS\b", r"\bCONCLUSION\b", r"\bRIGHT\s+TO\s+AMEND\b", r"\bAPPENDIX\b", r"\bANNEXURE\b", r"\bSECTION\s+20\b"],
    )
    return {"pre_accident_section_text": pre, "post_accident_section_text": post}


def _find_value(clean: str, labels: List[str], stop_labels: Optional[List[str]] = None) -> Optional[str]:
    stop_labels = stop_labels or ["Date", "Identity", "Gender", "Occupation", "Education", "Contact", "Age", "Current", "Attorney", "Reference"]
    for label in labels:
        pat = rf"(?i){label}\s*[:\-]?\s+(.{{1,120}}?)\s+(?:" + "|".join([re.escape(x) for x in stop_labels]) + r")\b"
        m = re.search(pat, clean)
        if m:
            return _norm_spaces(m.group(1))
        pat2 = rf"(?i){label}\s*[:\-]?\s+([^\n\r]{{1,120}})"
        m2 = re.search(pat2, clean)
        if m2:
            return _norm_spaces(m2.group(1))
    return None


def _extract_first(pattern: str, clean: str) -> Optional[str]:
    m = re.search(pattern, clean, flags=re.I)
    return _norm_spaces(m.group(1)) if m else None


def parse_base_case(clean: str) -> Dict[str, Any]:
    case: Dict[str, Any] = {}
    case["full_name"] = _find_value(clean, ["Full Name", "Claimant Name", "Name and Surname"])
    case["id_number"] = _find_value(clean, ["Identity Number", "ID Number"])
    case["dob_text"] = _find_value(clean, ["Date of Birth"])
    case["doa_text"] = _find_value(clean, ["Date of Accident", "Accident Date"])
    case["gender"] = _find_value(clean, ["Gender"])
    case["education"] = _find_value(clean, ["Education level", "Highest Qualification", "Highest Level of Education"])
    case["attorney_name"] = _find_value(clean, ["Referred By", "Attorney Name", "Instructing Attorneys"])
    case["attorney_reference"] = _find_value(clean, ["Your Reference", "Instruction Reference"])
    case["ip_report_date_text"] = _find_value(clean, ["Date of Report", "Report Release Date", "Report Date"])
    case["pre_occupation"] = _extract_first(r"(?i)(?:Pre\s*[-–]?\s*accident|Occupation at the time of the accident|PRE-ACCIDENT OCCUPATION)\s*[:]?\s*([^\n\r]{1,80})", clean)
    case["post_occupation"] = _extract_first(r"(?i)(?:Post\s*[-–]?\s*accident|Current Occupation|POST-ACCIDENT)\s*[:]?\s*([^\n\r]{1,80})", clean)
    if case.get("full_name"):
        parts = str(case["full_name"]).strip().split()
        if len(parts) > 1:
            case["surname"] = parts[-1]
            case["first_names"] = " ".join(parts[:-1])
    case.update(extract_earnings_sections(clean))
    return case


def parse_ip_report_ai(path: Path, use_llama: bool = True, llama_model: str = "llama3.1:8b") -> Dict[str, Any]:
    extracted = extract_ip_report_text(path)
    clean = extracted["clean"]
    case = parse_base_case(clean)
    case["extraction_method"] = extracted["method"]
    case["source_characters"] = extracted["characters"]
    from ai_extractor import ai_extract_ip_fields, map_ai_to_calculation_case
    ai = ai_extract_ip_fields(clean, use_llama=use_llama, llama_model=llama_model)
    return map_ai_to_calculation_case(case, ai)
