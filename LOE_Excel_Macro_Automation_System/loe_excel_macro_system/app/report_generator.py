
from __future__ import annotations

import shutil
import subprocess
from pathlib import Path
from typing import Any, Dict, Optional

from docx import Document
from docx.shared import Pt


def _get(case: Dict[str, Any], key: str, default: str = "Not specified") -> str:
    v = case.get(key) if isinstance(case, dict) else None
    return default if v in (None, "", [], {}) else str(v)


def _add_key_table(doc: Document, title: str, rows):
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        r = table.add_row().cells
        r[0].text = str(label)
        r[1].text = str(value or "Not specified")


def _add_summary_rows(doc: Document, title: str, rows):
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["Financial Year", "Age", "Salary", "Salary Today", "Financial Terms", "Progression"]
    for i, h in enumerate(headers): table.rows[0].cells[i].text = h
    for item in rows or []:
        r = table.add_row().cells
        r[0].text = str(item.get("fin_year") or item.get("financial_year") or "")
        r[1].text = str(item.get("age") or "")
        r[2].text = str(item.get("salary") or "")
        r[3].text = str(item.get("salary_today") or "")
        r[4].text = str(item.get("financial_terms") or "")
        r[5].text = str(item.get("progression") or "")


def generate_actuarial_report_docx(case: Dict[str, Any], earnings_description: str, output_path: Path) -> Path:
    output_path = Path(output_path); output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("ACTUARIAL LOSS OF EARNINGS REPORT", level=1)
    _add_key_table(doc, "Claimant Details", [
        ("Claimant", _get(case, "full_name")), ("Identity Number", _get(case, "id_number")), ("Date of Birth", _get(case, "dob_text")), ("Date of Accident", _get(case, "doa_text")), ("Attorney", _get(case, "attorney_name")), ("Reference", _get(case, "attorney_reference")), ("Industrial Psychologist", _get(case, "ip_name", "Industrial Psychologist")), ("IP Report Date", _get(case, "ip_report_date_text")),
    ])
    doc.add_heading("Earnings Descriptions", level=2)
    for line in (earnings_description or "").splitlines():
        line = line.strip()
        if not line: continue
        if line.lower().startswith(("pre-accident", "post-accident")):
            doc.add_heading(line, level=3)
        elif line.startswith("•"):
            doc.add_paragraph(line[1:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)
    _add_summary_rows(doc, "Pre-Accident Calculation Inputs", case.get("pre_rows") or [])
    _add_summary_rows(doc, "Post-Accident Calculation Inputs", case.get("post_rows") or [])
    doc.add_heading("Actuarial Note", level=2)
    doc.add_paragraph("The accompanying Excel calculation model has been populated using the actual earnings scenarios, financial terms, career progression, salary basis and retirement assumptions extracted from the Industrial Psychologist report. The red fields in the Summary and Report tabs are overwritten with IP report values and are not default assumptions.")
    doc.add_paragraph("Where the Industrial Psychologist report expresses earnings as Paterson bands, Koch quartiles, bargaining council rates, affidavit earnings, payslip values, bank-statement averages or other stated benchmarks, those references are preserved for actuarial review.")
    doc.save(str(output_path))
    return output_path


def try_convert_docx_to_pdf(docx_path: Path) -> Optional[Path]:
    docx_path = Path(docx_path)
    soffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not soffice: return None
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120)
        pdf = docx_path.with_suffix(".pdf")
        return pdf if pdf.exists() else None
    except Exception:
        return None
