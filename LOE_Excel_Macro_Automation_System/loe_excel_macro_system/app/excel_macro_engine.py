
from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from dateutil import parser as dtparser
from openpyxl import load_workbook


def parse_any_date(s: str):
    if not s: return None
    return dtparser.parse(str(s), dayfirst=True, fuzzy=True).date()


def _set_if(ws, cell: str, value: Any):
    if value not in (None, "", [], {}): ws[cell] = value


def _num_or_text(value: Any):
    if value in (None, "", [], {}): return None
    if isinstance(value, (int, float)): return value
    txt = str(value).strip()
    # leave benchmark scale phrases as text
    if any(k in txt.lower() for k in ["quartile", "median", "paterson", "bargaining", "minimum wage", "koch", "salaryexpert"]):
        return txt
    try:
        return float(txt.replace("R", "").replace("ZAR", "").replace(",", "").replace(" ", ""))
    except Exception:
        return txt


def _header_positions(ws) -> List[Tuple[int, int]]:
    positions = []
    for row in ws.iter_rows():
        vals = [str(c.value or "").strip().lower() for c in row]
        joined = "|".join(vals)
        if "financial year" in joined and "salary" in joined and "progression" in joined:
            # first cell in header row with financial year
            for c in row:
                if str(c.value or "").strip().lower() == "financial year":
                    positions.append((c.row, c.column))
                    break
    return positions


def _find_column_map(ws, header_row: int) -> Dict[str, int]:
    wanted = {
        "financial_year": ["financial year", "fin year"],
        "age": ["age"],
        "salary": ["salary"],
        "salary_today": ["salary today", "today"],
        "financial_terms": ["financial terms", "terms"],
        "progression": ["progression"],
    }
    colmap: Dict[str, int] = {}
    for cell in ws[header_row]:
        val = str(cell.value or "").strip().lower()
        for key, labels in wanted.items():
            if any(label == val or label in val for label in labels):
                colmap[key] = cell.column
    return colmap


def _clear_and_write_block(ws, header_row: int, rows: List[Dict[str, Any]], max_rows: int = 18):
    colmap = _find_column_map(ws, header_row)
    if not colmap:
        return False
    start = header_row + 1
    for r in range(start, start + max_rows):
        for key, col in colmap.items():
            # only clear values in input area; formulas in blank columns will be overwritten only if explicitly mapped.
            ws.cell(r, col).value = None
    for idx, item in enumerate(rows[:max_rows]):
        r = start + idx
        values = {
            "financial_year": item.get("fin_year") or item.get("financial_year"),
            "age": item.get("age"),
            "salary": _num_or_text(item.get("salary")),
            "salary_today": _num_or_text(item.get("salary_today")),
            "financial_terms": item.get("financial_terms"),
            "progression": item.get("progression") or "Flat",
        }
        for key, val in values.items():
            if key in colmap and val not in (None, "", [], {}):
                ws.cell(r, colmap[key]).value = val
    return True


def _write_summary_rows(summary, pre_rows: List[Dict[str, Any]], post_rows: List[Dict[str, Any]]):
    positions = _header_positions(summary)
    if positions:
        # first header block = pre, second header block = post where available.
        _clear_and_write_block(summary, positions[0][0], pre_rows, 20)
        if len(positions) > 1:
            _clear_and_write_block(summary, positions[1][0], post_rows, 20)
        else:
            # fallback to known post block if only one header detected
            _write_legacy_rows(summary, post_rows, 55, 20)
    else:
        _write_legacy_rows(summary, pre_rows, 5, 20)
        _write_legacy_rows(summary, post_rows, 55, 20)


def _write_legacy_rows(ws, rows: List[Dict[str, Any]], start_row: int, max_rows: int = 18):
    # C:Financial Year, D:Age, E:Salary, F:Salary Today, G:Financial terms, H:Progression
    for r in range(start_row, start_row + max_rows):
        for col in "CDEFGH":
            ws[f"{col}{r}"] = None
    for i, item in enumerate(rows[:max_rows]):
        r = start_row + i
        ws[f"C{r}"] = item.get("fin_year") or item.get("financial_year")
        ws[f"D{r}"] = item.get("age")
        ws[f"E{r}"] = _num_or_text(item.get("salary"))
        ws[f"F{r}"] = _num_or_text(item.get("salary_today"))
        ws[f"G{r}"] = item.get("financial_terms")
        ws[f"H{r}"] = item.get("progression") or "Flat"


def _write_earnings_descriptions(ws, case: Dict[str, Any], earnings_description: str):
    # Clear a reasonable area and write a downloadable-style narrative into the sheet.
    for r in range(1, 80):
        for c in range(1, 4):
            ws.cell(r, c).value = None
    ws["A1"] = "Earnings Description"
    row = 3
    for line in (earnings_description or "").splitlines():
        if line.strip():
            ws.cell(row, 1).value = line.strip()
            row += 1


def populate_template_openpyxl(xlsm_template: Path, output_xlsm: Path, case: dict, contingencies: dict, earnings_description: str = ""):
    xlsm_template, output_xlsm = Path(xlsm_template), Path(output_xlsm)
    output_xlsm.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(xlsm_template, output_xlsm)
    wb = load_workbook(output_xlsm, keep_vba=True)
    report = wb["Report"] if "Report" in wb.sheetnames else wb.active
    summary = wb["Summary"] if "Summary" in wb.sheetnames else None
    ed = wb["Earnings Descriptions"] if "Earnings Descriptions" in wb.sheetnames else None

    # Report tab actuals from IP report.
    _set_if(report, "B7", case.get("surname"))
    _set_if(report, "C7", case.get("first_names"))
    _set_if(report, "E7", str(case.get("id_number") or "").replace(" ", ""))
    for cell, key in [("F7", "dob_text"), ("H7", "doa_text"), ("C14", "instruction_date_text"), ("D21", "ip_report_date_text")]:
        if case.get(key):
            try: report[cell] = parse_any_date(case[key])
            except Exception: report[cell] = case[key]
    _set_if(report, "H11", case.get("accident_type") or "Motor Vehicle Accident")
    _set_if(report, "B11", case.get("attorney_name"))
    _set_if(report, "C12", case.get("attorney_reference"))
    _set_if(report, "B21", case.get("ip_name") or "Industrial Psychologist")
    report["C17"] = contingencies.get("past_pre", 0.05)
    report["D17"] = contingencies.get("past_post", 0.05)
    report["C18"] = contingencies.get("future_pre", 0.15)
    report["D18"] = contingencies.get("future_post", 0.25)

    if summary:
        # Peak/retirement ages if template uses these red cells.
        if case.get("pre_peak_age") not in (None, ""): summary["C23"] = case.get("pre_peak_age")
        if case.get("pre_ret_age") not in (None, ""): summary["H23"] = case.get("pre_ret_age")
        if case.get("post_peak_age") not in (None, ""): summary["C73"] = case.get("post_peak_age")
        if case.get("post_ret_age") not in (None, ""): summary["H73"] = case.get("post_ret_age")
        _write_summary_rows(summary, case.get("pre_rows") or [], case.get("post_rows") or [])
        try: summary["C113"] = "N"
        except Exception: pass

    if ed:
        _write_earnings_descriptions(ed, case, earnings_description)

    wb.save(output_xlsm)
    return output_xlsm


def run_excel_macro(output_xlsm: Path, docm_template_dir: Path, save_root: Path):
    import win32com.client as win32
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AskToUpdateLinks = False
    wb = None
    try:
        wb = excel.Workbooks.Open(str(output_xlsm))
        ws = wb.Worksheets("Report")
        ws.Range("C2").Value = str(docm_template_dir)
        ws.Range("C3").Value = str(save_root)
        excel.CalculateFullRebuild()
        wb.Save()
        excel.Application.Run(f"'{wb.Name}'!LOEOpenWordandSave")
        wb.Save()
    finally:
        if wb is not None:
            wb.Close(SaveChanges=True)
        excel.Quit()


def export_latest_docm_to_pdf(search_root: Path):
    import win32com.client as win32
    docm_files = sorted(Path(search_root).rglob("*.docm"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not docm_files: return None
    src = docm_files[0]
    pdf_path = src.with_suffix(".pdf")
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    doc = None
    try:
        doc = word.Documents.Open(str(src))
        doc.SaveAs(str(pdf_path), FileFormat=17)
    finally:
        if doc is not None: doc.Close(False)
        word.Quit()
    return pdf_path
